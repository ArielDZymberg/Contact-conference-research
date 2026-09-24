#!/usr/bin/env python3
"""Adiciona briefings à base de contatos em Excel (cria a base se não existir).

Uso:
  python base_excel.py --briefings briefings.json --data-conversa 2026-09-24 \
      --conferencia "Conferência X" [--observacoes "..."] [--base base.xlsx] --saida base.xlsx
"""

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET = "Contatos"
COLUMNS = [
    ("Data da conversa", 16),
    ("Conferência", 28),
    ("Nome", 26),
    ("Cargo", 28),
    ("Fundo/Empresa", 28),
    ("Tipo", 18),
    ("Estratégia", 20),
    ("AUM", 18),
    ("AUM - data de referência", 16),
    ("Resumo da pessoa", 60),
    ("Resumo do fundo/empresa", 60),
    ("LinkedIn", 36),
    ("Site", 30),
    ("Observações", 40),
    ("Adicionado em", 18),
]
HEADERS = [c[0] for c in COLUMNS]


def new_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADERS)
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    return wb


def open_base(path):
    if not path:
        return new_workbook()
    wb = load_workbook(path)
    if SHEET not in wb.sheetnames:
        # Planilha de outro formato: usa a primeira aba se o cabeçalho bater, senão cria a aba.
        first = wb.worksheets[0]
        header = [c.value for c in first[1]]
        if header[: len(HEADERS)] == HEADERS:
            first.title = SHEET
        else:
            ws = wb.create_sheet(SHEET)
            ws.append(HEADERS)
    return wb


def as_key(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value or "").strip().lower()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--briefings", required=True, help="JSON com lista de briefings (ou um único objeto)")
    ap.add_argument("--data-conversa", required=True, help="AAAA-MM-DD")
    ap.add_argument("--conferencia", required=True)
    ap.add_argument("--observacoes", default="")
    ap.add_argument("--base", help="Base existente (.xlsx). Omitir para criar uma nova.")
    ap.add_argument("--saida", required=True)
    args = ap.parse_args()

    conv_date = date.fromisoformat(args.data_conversa)
    data = json.loads(Path(args.briefings).read_text(encoding="utf-8"))
    briefings = data if isinstance(data, list) else [data]

    wb = open_base(args.base)
    ws = wb[SHEET]
    existing = {
        (as_key(r[0]), as_key(r[1]), as_key(r[2]))
        for r in ws.iter_rows(min_row=2, max_col=3, values_only=True)
        if any(r)
    }

    added, skipped = [], []
    for b in briefings:
        pessoa, empresa = b.get("pessoa", {}), b.get("empresa", {})
        estrategia, aum = b.get("estrategia", {}), b.get("aum", {})
        key = (as_key(conv_date), as_key(args.conferencia), as_key(pessoa.get("nome")))
        if key in existing:
            skipped.append(pessoa.get("nome", ""))
            continue
        row = [
            conv_date,
            args.conferencia.strip(),
            pessoa.get("nome", ""),
            pessoa.get("cargo", ""),
            empresa.get("nome", ""),
            empresa.get("tipo", ""),
            estrategia.get("tipo", "") if estrategia.get("aplicavel") else "N/A",
            aum.get("valor", ""),
            aum.get("data_referencia", ""),
            pessoa.get("resumo", ""),
            empresa.get("resumo", ""),
            pessoa.get("linkedin_url", ""),
            empresa.get("site", ""),
            args.observacoes.strip(),
            datetime.now().replace(microsecond=0),
        ]
        ws.append(row)
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=15).number_format = "DD/MM/YYYY HH:MM"
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        existing.add(key)
        added.append(pessoa.get("nome", ""))

    Path(args.saida).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.saida)
    total = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if any(r))
    print(f"Adicionados: {', '.join(added) or 'nenhum'}")
    if skipped:
        print(f"Já estavam na base (mesma data e conferência), não duplicados: {', '.join(skipped)}")
    print(f"Total de contatos na base: {total}")
    print(f"Arquivo salvo em: {args.saida}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
