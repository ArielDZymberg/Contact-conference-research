"""Extração de participantes (nome + fundo/empresa) a partir da agenda da conferência."""

import base64
import csv
import io
import json
from pathlib import Path

from research import FALLBACK_BETA, MODEL, ResearchError, get_client

SUPPORTED = {".pdf", ".xlsx", ".xlsm", ".csv", ".txt", ".md", ".docx"}

EXTRACT_SYSTEM = """Você extrai a lista de pessoas de uma agenda de conferência. Para cada pessoa, traga o nome \
completo, o fundo/empresa e, se houver, o cargo e o painel/reunião/horário em que aparece. Não repita a mesma \
pessoa. Ignore moderadores genéricos sem nome, nomes de salas e patrocinadores sem pessoa associada. Use string \
vazia quando um campo não constar na agenda."""

PEOPLE_SCHEMA = {
    "type": "object",
    "properties": {
        "pessoas": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "nome": {"type": "string"},
                    "empresa": {"type": "string"},
                    "cargo": {"type": "string"},
                    "sessao": {"type": "string", "description": "Painel, reunião ou horário"},
                },
                "required": ["nome", "empresa", "cargo", "sessao"],
                "additionalProperties": False,
            },
        }
    },
    "required": ["pessoas"],
    "additionalProperties": False,
}


def _xlsx_to_text(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# Planilha: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _csv_to_text(data: bytes) -> str:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text), dialect) if any(row))


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def file_to_content(filename: str, data: bytes) -> list[dict]:
    """Converte o arquivo em blocos de conteúdo para enviar ao Claude."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ResearchError(
            f"Formato '{ext or filename}' não suportado. Use PDF, Excel (.xlsx), CSV, Word (.docx) ou texto."
        )
    if ext == ".pdf":
        return [
            {
                "type": "document",
                "source": {
                    "type": "base64",
                    "media_type": "application/pdf",
                    "data": base64.standard_b64encode(data).decode(),
                },
            }
        ]
    if ext in (".xlsx", ".xlsm"):
        text = _xlsx_to_text(data)
    elif ext == ".csv":
        text = _csv_to_text(data)
    elif ext == ".docx":
        text = _docx_to_text(data)
    else:
        text = data.decode("utf-8", errors="replace")
    if not text.strip():
        raise ResearchError("O arquivo parece estar vazio.")
    return [{"type": "text", "text": f"<agenda>\n{text}\n</agenda>"}]


def extract_people(filename: str, data: bytes) -> list[dict]:
    content = file_to_content(filename, data)
    content.append({"type": "text", "text": "Extraia todas as pessoas desta agenda."})
    with get_client().beta.messages.stream(
        model=MODEL,
        max_tokens=64000,
        system=EXTRACT_SYSTEM,
        output_config={"effort": "low", "format": {"type": "json_schema", "schema": PEOPLE_SCHEMA}},
        messages=[{"role": "user", "content": content}],
        betas=[FALLBACK_BETA],
        fallbacks="default",
    ) as stream:
        response = stream.get_final_message()
    if response.stop_reason == "refusal":
        raise ResearchError("O modelo recusou processar este arquivo.")
    if response.stop_reason == "max_tokens":
        raise ResearchError("A agenda é grande demais para ser processada de uma vez. Divida o arquivo em partes.")
    text = next((b.text for b in response.content if b.type == "text"), "")
    people = json.loads(text)["pessoas"] if text else []
    return [p for p in people if p["nome"].strip()]
