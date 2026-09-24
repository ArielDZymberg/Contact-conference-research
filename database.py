"""Base histórica de contatos em Excel."""

import os
import threading
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA_DIR = Path(os.environ.get("DATA_DIR", Path(__file__).parent / "dados"))
DB_PATH = DATA_DIR / "base_contatos.xlsx"
SHEET = "Contatos"

# (cabeçalho, largura da coluna)
COLUMNS = [
    ("Data da conversa", 16),
    ("Conferência", 28),
    ("Nome", 26),
    ("Cargo", 28),
    ("Fundo/Empresa", 28),
    ("Tipo", 18),
    ("Estratégia", 20),
    ("AUM", 18),
    ("AUM - data de referência", 16),
    ("Resumo da pessoa", 60),
    ("Resumo do fundo/empresa", 60),
    ("LinkedIn", 36),
    ("Site", 30),
    ("Observações", 40),
    ("Adicionado em", 18),
]

_lock = threading.Lock()


class DatabaseError(Exception):
    pass


def _create() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append([c[0] for c in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    return wb


def _open() -> Workbook:
    if DB_PATH.exists():
        return load_workbook(DB_PATH)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    return _create()


def _save(wb: Workbook) -> None:
    try:
        wb.save(DB_PATH)
    except PermissionError as e:
        raise DatabaseError(
            "Não foi possível salvar a base. Feche o arquivo base_contatos.xlsx no Excel e tente novamente."
        ) from e


def add_contact(result: dict, conversation_date: str, conference: str, notes: str = "") -> int:
    """Adiciona uma linha à base e devolve o total de contatos."""
    try:
        conv_date = date.fromisoformat(conversation_date)
    except (TypeError, ValueError) as e:
        raise DatabaseError("Data da conversa inválida.") from e
    if not conference.strip():
        raise DatabaseError("Informe a conferência.")

    pessoa, empresa = result.get("pessoa", {}), result.get("empresa", {})
    estrategia, aum = result.get("estrategia", {}), result.get("aum", {})
    row = [
        conv_date,
        conference.strip(),
        pessoa.get("nome", ""),
        pessoa.get("cargo", ""),
        empresa.get("nome", ""),
        empresa.get("tipo", ""),
        estrategia.get("tipo", "") if estrategia.get("aplicavel") else "N/A",
        aum.get("valor", ""),
        aum.get("data_referencia", ""),
        pessoa.get("resumo", ""),
        empresa.get("resumo", ""),
        pessoa.get("linkedin_url", ""),
        empresa.get("site", ""),
        notes.strip(),
        datetime.now().replace(microsecond=0),
    ]
    with _lock:
        wb = _open()
        ws = wb[SHEET]
        ws.append(row)
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=15).number_format = "DD/MM/YYYY HH:MM"
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="top", wrap_text=True)
        _save(wb)
        return r - 1


def list_contacts() -> list[dict]:
    if not DB_PATH.exists():
        return []
    with _lock:
        wb = load_workbook(DB_PATH, read_only=True, data_only=True)
        rows = list(wb[SHEET].iter_rows(min_row=2, values_only=True))
        wb.close()
    headers = [c[0] for c in COLUMNS]
    out = []
    for row in rows:
        if not any(row):
            continue
        item = {}
        for h, v in zip(headers, row):
            if isinstance(v, datetime):
                v = v.strftime("%d/%m/%Y %H:%M") if h == "Adicionado em" else v.strftime("%d/%m/%Y")
            elif isinstance(v, date):
                v = v.strftime("%d/%m/%Y")
            item[h] = "" if v is None else str(v)
        out.append(item)
    return out


def conferences() -> list[str]:
    seen = []
    for c in list_contacts():
        name = c.get("Conferência", "")
        if name and name not in seen:
            seen.append(name)
    return seen
